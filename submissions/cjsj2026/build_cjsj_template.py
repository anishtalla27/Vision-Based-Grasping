"""Fill the official CJSJ Original Research template with the paper text."""
import copy, re, zipfile, shutil, os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import cjsj_text as T

TEMPLATE = "official/CJSJ-Original-Research-Template-1-faxy.docx"
OUT = "TallaAnish_paper.docx"
doc = Document(TEMPLATE)
body = doc.element.body
P = list(doc.paragraphs)
proto = dict(spacer=P[0]._p, title=P[1]._p, authors=P[2]._p, abstract=P[13]._p, heading=P[15]._p,
             text=P[16]._p, tabtitle=P[24]._p, figcap=P[26]._p, ackhead=P[28]._p, refhead=P[30]._p, ref=P[31]._p)
proto = {k: copy.deepcopy(v) for k, v in proto.items()}
sect = body.find(qn("w:sectPr"))
for el in list(body):
    if el is not sect: body.remove(el)

def W(tag): return qn("w:" + tag)
def make(kind, text, keep_footnote=False):
    p = copy.deepcopy(proto[kind])
    runs = p.findall(W("r"))
    base = None
    for r in runs:
        if r.find(W("t")) is not None and (r.find(W("t")).text or "").strip():
            base = r; break
    if base is None:
        withs = [r for r in runs if r.find(W("t")) is not None and r.find(W("footnoteReference")) is None]
        base = withs[0] if withs else runs[0]
    if base.find(W("t")) is None:
        base.append(base.makeelement(W("t"), {}))
    fn = [r for r in runs if r.find(W("footnoteReference")) is not None]
    for r in runs:
        if r is not base and not (keep_footnote and r in fn): p.remove(r)
    for child in list(p):
        if child.tag in (W("hyperlink"), W("bookmarkStart"), W("bookmarkEnd")): p.remove(child)
    for t in base.findall(W("t"))[1:]: base.remove(t)
    for extra in base.findall(W("tab")) + base.findall(W("br")): base.remove(extra)
    t = base.find(W("t")); t.text = text; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    if keep_footnote and fn:  # footnote mark goes after the text
        for r in fn: p.remove(r); p.append(r)
    sect.addprevious(p)
    return p

def lead_italic(p, lead):
    """Put an italic run-in subheading in front of a body paragraph."""
    base = [r for r in p.findall(W("r")) if r.find(W("t")) is not None][0]
    r2 = copy.deepcopy(base); r2.find(W("t")).text = lead + ". "
    rpr = r2.find(W("rPr"))
    for tag in ("i", "iCs"):
        e = rpr.find(W(tag))
        if e is None: e = rpr.makeelement(W(tag), {}); rpr.append(e)
        e.set(W("val"), "1")
    base.addprevious(r2)

def add_table(title, header, rows):
    make("tabtitle", title)
    tbl = doc.add_table(rows=1, cols=len(header)); tbl.autofit = False
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    borders = parse_xml('<w:tblBorders %s>' % nsdecls("w") + "".join(
        '<w:%s w:val="single" w:sz="4" w:space="0" w:color="000000"/>' % e for e in ("top","left","bottom","right","insideH","insideV")) + '</w:tblBorders>')
    tbl._tbl.tblPr.append(borders)
    for i, h in enumerate(header):
        c = tbl.rows[0].cells[i]; c.text = h
    for row in rows:
        cells = tbl.add_row().cells
        for i, v in enumerate(row): cells[i].text = v
    for ri, row in enumerate(tbl.rows):
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0); par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.0; par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in par.runs:
                    r.font.size = Pt(8); r.font.name = "Times New Roman"; r.bold = (ri == 0)
    widths = [1.45, 0.85, 1.0] if len(header) == 3 else [0.98, 0.52, 0.48, 0.54, 0.78]
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.oxml import OxmlElement
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tbl._tbl.tblPr.append(lay)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 1440)))
    for row in tbl.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for c, w in zip(row.cells, widths): c.width = Inches(w)
    sect.addprevious(tbl._tbl)
    spacer = make("text", " ")

def add_figure(img, caption):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(img, width=Inches(3.3))
    par.paragraph_format.space_before = Pt(6)
    sect.addprevious(par._p)
    make("figcap", caption)

make("spacer", " ", keep_footnote=True)
make("title", T.TITLE, keep_footnote=True)
make("authors", T.AUTHOR, keep_footnote=True)
make("abstract", "Abstract—" + T.ABSTRACT)

figs = list(T.CAPS.items())
for name, paras in T.SECTIONS:
    make("heading", name)
    for i, item in enumerate(paras):
        if isinstance(item, tuple):
            p = make("text", item[1]); lead_italic(p, item[0])
        else:
            make("text", item)
        if name.startswith("Results"):
            if i == 0: add_table(*T.TABLE1); add_figure(*figs[0])
            if i == 1: add_figure(*figs[1])
            if i == 2: add_table(*T.TABLE2); add_figure(*figs[2])
make("ackhead", "Acknowledgment")
make("text", T.ACK)
make("refhead", "References")
for r in T.REFS: make("ref", r)

# Appendix: exact prompts (CJSJ asks for full prompts when AI is part of the experiment)
src = open("TallaAnish_supplement_prompts.md").read()
blocks = re.findall(r"### (.*?)\n\n```\n(.*?)\n```", src, re.S)
make("refhead", "Appendix: Prompts Given to GPT-4o")
make("text", "Model openai/gpt-4o through OpenRouter, default temperature, no seed. In the System D prompt, {k} is the number of candidates drawn. The version with every mark the same size adds one sentence to the System D prompt, shown last.")
labels = ["System C, system message", "System C, user prompt", "System D, system message", "System D, user prompt"]
for (h, txt), lab in zip(blocks[:4], labels):
    txt = " ".join(txt.replace("{{", "{").replace("}}", "}").split())
    p = make("ref", txt); 
    # unnumbered small text: drop list numbering
    ppr = p.find(W("pPr")); n = ppr.find(W("numPr")) if ppr is not None else None
    if n is not None: ppr.remove(n)
    lead_italic(p, lab)
extra = "All candidates are drawn at the same length, which is NOT to scale: a mark shows only where the grip is centred and the direction the fingertips close along. So a candidate grips the part of the object that the line between its two bars crosses, at the marked centre."
p = make("ref", extra); ppr = p.find(W("pPr")); n = ppr.find(W("numPr"))
if n is not None: ppr.remove(n)
lead_italic(p, "Sentence added for the same-size version")

doc.save(OUT)

# footnote text: replace template's author footnotes with one line
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/footnotes.xml":
            from lxml import etree
            NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            root = etree.fromstring(data)
            fn = [f for f in root.findall("{%s}footnote" % NS) if f.get("{%s}id" % NS) == "0"][0]
            paras = fn.findall("{%s}p" % NS)
            for extra in paras[1:]: fn.remove(extra)
            first = True
            for t in paras[0].iter("{%s}t" % NS):
                t.text = T.FOOTNOTE if first else ""
                first = False
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(item, data)
os.replace(tmp, OUT)
print("saved", OUT)
