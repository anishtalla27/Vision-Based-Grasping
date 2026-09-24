"""Fill the official CJSJ figures template."""
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
import cjsj_text as T
prs = Presentation("official/CJSJ-Figures-Template.pptx")
s0, sfig, stab = prs.slides[0], prs.slides[1], prs.slides[2]
ph = [sh for sh in s0.shapes if sh.has_text_frame]
ph[0].text_frame.text = "Three Ways to Miss a Grasp: Figures and Tables"
ph[1].text_frame.text = "Anish Talla, 2026-2027"
def fill(slide, img, caption, maxw=6.75, maxh=4.5):
    for p in [sh for sh in slide.shapes if sh.shape_type == 13]: p._element.getparent().remove(p._element)
    w, h = Image.open(img).size; sc = min(maxw / w, maxh / h); W_, H_ = w * sc, h * sc
    slide.shapes.add_picture(img, Inches((10 - W_) / 2), Inches(0.13 + (4.62 - H_) / 2), width=Inches(W_))
    cap = [sh for sh in slide.shapes if sh.has_text_frame][0]
    cap.left, cap.top, cap.width, cap.height = Inches(0.34), Inches(4.83), Inches(9.32), Inches(0.63)
    cap.text_frame.text = caption
    for para in cap.text_frame.paragraphs:
        for r in para.runs: r.font.size = Pt(12)
def new_like(src, img, cap, **kw):
    s = prs.slides.add_slide(src.slide_layout); fill(s, img, cap, **kw); return s
items = list(T.CAPS.items())
fill(sfig, items[0][0], "Figure 1.\t" + items[0][1])
fill(stab, "table-1.png", "Table 1.\t" + T.TABLE1[0] + ".", maxw=7.5, maxh=3.2)
new_like(sfig, items[1][0], "Figure 2.\t" + items[1][1])
new_like(sfig, items[2][0], "Figure 3.\t" + items[2][1], maxw=5.2)
new_like(stab, "table-2.png", "Table 2.\t" + T.TABLE2[0] + ".", maxw=7.5, maxh=2.6)
# order: title, Fig 1-3, Table 1-2
lst = prs.slides._sldIdLst; ids = list(lst)
for el in ids: lst.remove(el)
for i in [0, 1, 3, 4, 2, 5]: lst.append(ids[i])
prs.save("TallaAnish_figures.pptx"); print("pptx saved", len(prs.slides))
